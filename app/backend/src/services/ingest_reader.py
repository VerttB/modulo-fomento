import csv
import io
import re
import unicodedata
import zipfile
import xml.etree.ElementTree as ET
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl
import xlrd
from openpyxl.utils.datetime import from_excel

from src.core.exceptions import IngestFormatError


def normalizar(texto: str) -> str:
    texto = unicodedata.normalize("NFKD", texto.casefold())
    texto = "".join(char for char in texto if not unicodedata.combining(char))
    return " ".join(re.findall(r"[a-z0-9]+", texto))


def ler_csv(conteudo: bytes | Path) -> list[list[Any]]:
    if isinstance(conteudo, Path):
        with conteudo.open("rb") as stream:
            amostra = stream.read(8192)
        try:
            amostra.decode("utf-8-sig")
            encoding = "utf-8-sig"
        except UnicodeDecodeError:
            encoding = "cp1252"
        with conteudo.open("r", encoding=encoding, newline="") as stream:
            amostra_texto = stream.read(8192)
            try:
                delimiter = csv.Sniffer().sniff(
                    amostra_texto, delimiters=";,\t"
                ).delimiter
            except csv.Error:
                delimiter = ";"
            stream.seek(0)
            return [list(row) for row in csv.reader(stream, delimiter=delimiter)]
    try:
        texto = conteudo.decode("utf-8-sig")
    except UnicodeDecodeError:
        texto = conteudo.decode("cp1252")
    try:
        delimitador = csv.Sniffer().sniff(texto[:8192], delimiters=";,\t").delimiter
    except csv.Error:
        delimitador = ";"
    return [list(row) for row in csv.reader(io.StringIO(texto), delimiter=delimitador)]


def _ler_ods(conteudo: bytes | Path) -> list[list[Any]]:
    namespaces = {
        "table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0",
        "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
        "office": "urn:oasis:names:tc:opendocument:xmlns:office:1.0",
    }
    rows: list[list[Any]] = []
    try:
        source = conteudo if isinstance(conteudo, Path) else io.BytesIO(conteudo)
        with zipfile.ZipFile(source) as document:
            with document.open("content.xml") as content:
                for _, element in ET.iterparse(content, events=("end",)):
                    if element.tag != f"{{{namespaces['table']}}}table-row":
                        continue
                    row: list[Any] = []
                    for cell in list(element):
                        if cell.tag not in {
                            f"{{{namespaces['table']}}}table-cell",
                            f"{{{namespaces['table']}}}covered-table-cell",
                        }:
                            continue
                        value = cell.get(f"{{{namespaces['office']}}}date-value")
                        if value:
                            parsed: Any = value
                        else:
                            value = cell.get(f"{{{namespaces['office']}}}value")
                            value_type = cell.get(f"{{{namespaces['office']}}}value-type")
                            if value is not None and value_type in {"float", "currency", "percentage"}:
                                try:
                                    parsed = float(value)
                                except ValueError:
                                    parsed = value
                            elif value_type == "boolean":
                                parsed = cell.get(f"{{{namespaces['office']}}}boolean-value") == "true"
                            else:
                                paragraphs = cell.findall(f".//{{{namespaces['text']}}}p")
                                parsed = "".join("".join(paragraph.itertext()) for paragraph in paragraphs)
                        repeat = int(cell.get(f"{{{namespaces['table']}}}number-columns-repeated", "1"))
                        row.extend([parsed] * min(repeat, 256))
                    while row and row[-1] == "":
                        row.pop()
                    row_repeat = int(element.get(f"{{{namespaces['table']}}}number-rows-repeated", "1"))
                    rows.extend([row.copy() for _ in range(min(row_repeat, 256))])
                    element.clear()
    except (KeyError, zipfile.BadZipFile, ET.ParseError, ValueError) as error:
        raise IngestFormatError("Não foi possível ler o arquivo ODS enviado.") from error
    return rows


def ler_planilha(
    nome_arquivo: str, conteudo: bytes | Path
) -> tuple[list[list[Any]], Any]:
    extensao = Path(nome_arquivo).suffix.casefold()
    if extensao == ".csv":
        return ler_csv(conteudo), None
    if extensao == ".xlsx":
        source = conteudo if isinstance(conteudo, Path) else io.BytesIO(conteudo)
        workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
        return [list(row) for row in workbook.active.iter_rows(values_only=True)], workbook.epoch
    if extensao == ".xls":
        workbook = (
            xlrd.open_workbook(filename=str(conteudo))
            if isinstance(conteudo, Path)
            else xlrd.open_workbook(file_contents=conteudo)
        )
        sheet = workbook.sheet_by_index(0)
        rows: list[list[Any]] = []
        for row_index in range(sheet.nrows):
            row: list[Any] = []
            for column_index in range(sheet.ncols):
                cell = sheet.cell(row_index, column_index)
                value: Any = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    value = xlrd.xldate_as_datetime(value, workbook.datemode)
                elif cell.ctype == xlrd.XL_CELL_EMPTY:
                    value = None
                row.append(value)
            rows.append(row)
        return rows, None
    if extensao == ".ods":
        return _ler_ods(conteudo), None
    raise IngestFormatError("Formato inválido. Envie um arquivo CSV, XLS, XLSX ou ODS.")


def mapear_colunas(cabecalho: list[Any]) -> dict[str, int]:
    indices: dict[str, int] = {}
    finais = 0
    for indice, valor in enumerate(cabecalho):
        nome = normalizar(str(valor or ""))
        if nome == "data final" or nome.startswith("data final "):
            finais += 1
            indices[f"data final {finais}"] = indice
        elif nome.startswith("pal chave"):
            sufixo = nome.removeprefix("pal chave")
            if sufixo in {"1", "2", "3", "4"}:
                indices[f"pal chave{sufixo}"] = indice
        else:
            indices[nome] = indice

    # normalizar("cpf_pesquisador") resulta em "cpf pesquisador".
    # O ingest procura pelo campo cpf_pesquisador (aceitando variações
    # como "CPF Pesquisador") e mantém compatibilidade com "cpf" legado.
    if "cpf pesquisador" in indices:
        indices["cpf_pesquisador"] = indices["cpf pesquisador"]
    elif "cpf" in indices:
        indices["cpf_pesquisador"] = indices["cpf"]

    # normalizar("sub área") / normalizar("sub-área") resulta em "sub area",
    # enquanto normalizar("subárea") resulta em "subarea".
    # O ingest aceita com ou sem acento, com ou sem espaço/hífen.
    if "subarea" in indices:
        indices["sub area"] = indices["subarea"]
    elif "sub area" in indices:
        indices["subarea"] = indices["sub area"]

    # Único cabeçalho obrigatório do ingest: cpf_pesquisador.
    # Qualquer outro campo pode vir vazio (ou até sem coluna) e resulta em None.
    obrigatorias = ("cpf_pesquisador",)
    ausentes = [nome for nome in obrigatorias if nome not in indices]
    if ausentes:
        raise IngestFormatError(
            "Cabeçalhos obrigatórios ausentes: " + ", ".join(ausentes)
            + ". Inclua uma coluna cpf_pesquisador para identificar cada pesquisador."
        )
    return indices


def texto_linha(row: list[Any], indices: dict[str, int], coluna: str, obrigatorio: bool = False) -> str | None:
    indice = indices.get(coluna)
    valor = row[indice] if indice is not None and indice < len(row) else None
    if valor is None:
        texto = ""
    elif isinstance(valor, float) and valor.is_integer():
        texto = str(int(valor))
    else:
        texto = str(valor).strip()
    if not texto:
        if obrigatorio:
            raise IngestFormatError(f"Campo obrigatório vazio: {coluna}.")
        return None
    return texto


def normalizar_cpf(valor: str) -> str:
    digitos = re.sub(r"\D", "", valor)
    if len(digitos) > 11:
        raise IngestFormatError("CPF inválido: esperado um CPF com 11 dígitos.")
    digitos = digitos.zfill(11)
    if len(digitos) != 11 or len(set(digitos)) == 1:
        raise IngestFormatError("CPF inválido.")
    return digitos


def parse_data(valor: Any, epoch: Any = None) -> date | None:
    if valor is None or str(valor).strip() == "":
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, (int, float)) and epoch is not None:
        return from_excel(valor, epoch).date()
    texto = str(valor).strip()
    try:
        return datetime.fromisoformat(texto).date()
    except ValueError:
        pass
    for formato in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            continue
    raise IngestFormatError(f"Data inválida: {texto!r}.")
